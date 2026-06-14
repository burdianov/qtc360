"""Excel-based checklist PDF generation.

Fills an XLSX template with checklist data (placeholders like {{ sn_1 }}, {{ item_1 }},
{{ yes_1 }}, {{ no_1 }}, {{ na_1 }}, {{ display_name }}) then converts to PDF via Gotenberg.
"""

import io
import re

import httpx
from openpyxl import load_workbook

from app.core.config import settings
from app.core.types import GOTENBERG_TIMEOUT

PLACEHOLDER_RE = re.compile(r"\{\{\s*(\w+)\s*\}\}")
TICK = "\u2713"  # ✓ (check mark)


def fill_xlsx_template(template_bytes: bytes, context: dict[str, str]) -> bytes:
    """Replace {{ placeholder }} patterns in all cells of an XLSX workbook."""
    from openpyxl.styles import Font

    wb = load_workbook(io.BytesIO(template_bytes))
    for ws in wb.worksheets:
        # Capture the font from a tick placeholder cell ({{ yes_N }}) —
        # the tick mark will use this so it matches the placeholder text
        # that the Excel template author designed.
        placeholder_font: Font | None = None
        for row in ws.iter_rows():
            for cell in row:
                if cell.value and isinstance(cell.value, str) and "{{" in cell.value:
                    # Capture from the first yes/no placeholder cell
                    if placeholder_font is None and ("{{ yes_" in cell.value or "{{ no_" in cell.value):
                        placeholder_font = Font(
                            name=cell.font.name or "Arial",
                            size=cell.font.size or 10,
                            bold=cell.font.bold or False,
                            color=cell.font.color,
                        )
                    new_value = PLACEHOLDER_RE.sub(
                        lambda m: context.get(m.group(1), ""), cell.value
                    )
                    cell.value = new_value
                    # Apply the placeholder font to the tick mark
                    if new_value == TICK and placeholder_font is not None:
                        cell.font = placeholder_font
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def convert_xlsx_to_pdf(xlsx_bytes: bytes) -> bytes:
    """Convert XLSX to PDF via Gotenberg's LibreOffice route."""
    url = f"{settings.gotenberg_url.rstrip('/')}/forms/libreoffice/convert"
    async with httpx.AsyncClient(timeout=GOTENBERG_TIMEOUT) as client:
        resp = await client.post(
            url,
            files={
                "files": (
                    "spreadsheet.xlsx",
                    xlsx_bytes,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
        if resp.status_code != 200:
            raise RuntimeError(
                f"Gotenberg XLSX->PDF conversion failed: {resp.status_code}"
            )
        return resp.content


def build_checklist_context(
    display_name: str,
    responses: list[dict],
) -> dict[str, str]:
    """Build the placeholder context dict for a checklist.

    responses: list of {item_text, response, notes} dicts ordered by display_order.
    response values: 'yes', 'no', or '' (unfilled).
    notes: free-text notes (shown in the Notes column).
    """
    ctx: dict[str, str] = {"display_name": display_name}
    for i in range(1, 21):
        if i <= len(responses):
            r = responses[i - 1]
            ctx[f"sn_{i}"] = str(i)
            ctx[f"item_{i}"] = r["item_text"]
            ctx[f"yes_{i}"] = TICK if r["response"] == "yes" else ""
            ctx[f"no_{i}"] = TICK if r["response"] == "no" else ""
            ctx[f"na_{i}"] = r.get("notes", "") or ""
        else:
            ctx[f"sn_{i}"] = ""
            ctx[f"item_{i}"] = ""
            ctx[f"yes_{i}"] = ""
            ctx[f"no_{i}"] = ""
            ctx[f"na_{i}"] = ""
    return ctx
